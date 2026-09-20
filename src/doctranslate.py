"""DocTranslate En->Vi: docx/txt -> docx/txt (kieu Saydi). Dung NMT + ContextEngine hien co."""
import os

def translate_docx(in_path, out_path, translator, context=None):
    from docx import Document
    doc = Document(in_path)
    for p in doc.paragraphs:
        if p.text.strip():
            vi, *_ = (context.translate(p.text, translator) if context
                       else (translator.translate(p.text)[0],))
            if isinstance(vi, tuple):
                vi = vi[0]
            p.text = vi if isinstance(vi, str) else str(vi)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.text.strip():
                        vi, *_ = (context.translate(p.text, translator) if context
                                   else (translator.translate(p.text)[0],))
                        if isinstance(vi, tuple):
                            vi = vi[0]
                        p.text = vi if isinstance(vi, str) else str(vi)
    doc.save(out_path)
    return out_path

def translate_txt(in_path, out_path, translator, context=None):
    with open(in_path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    out = []
    for ln in lines:
        if ln.strip():
            vi, *_ = (context.translate(ln, translator) if context
                       else (translator.translate(ln)[0],))
            if isinstance(vi, tuple):
                vi = vi[0]
            out.append(vi if isinstance(vi, str) else str(vi))
        else:
            out.append("")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    return out_path

def translate_doc(in_path, out_path=None, translator=None, context=None):
    if translator is None:
        from .nmt import NMTTranslator
        translator = NMTTranslator()
    if out_path is None:
        base, ext = os.path.splitext(in_path)
        out_path = base + ".vi" + ext
    ext = os.path.splitext(in_path)[1].lower()
    if ext == ".docx":
        return translate_docx(in_path, out_path, translator, context)
    return translate_txt(in_path, out_path, translator, context)
