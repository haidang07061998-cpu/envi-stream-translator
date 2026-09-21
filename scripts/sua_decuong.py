"""Sua DeCuong.docx theo dung nhung gi da lam: mo hinh nao, du lieu nao.

- Giu NGUYEN form cua truong (header, muc 1-2, bang tien do, khoi chu ky).
- Thay noi dung muc 3 -> 7 bang ban da cap nhat.
- Luu ra file MOI, khong ghi de ban goc.
"""
import copy
import os
import sys

import docx
from docx.text.paragraph import Paragraph

SRC = "DeCuong - NguyenHoangKhoiNguyen.docx"
DST = "DeCuong - NguyenHoangKhoiNguyen (da sua).docx"

NOI_DUNG = [
    "3. Tính cấp thiết:",
    "Trong bối cảnh toàn cầu hóa và chuyển đổi số, nhu cầu giao tiếp, hội họp, học tập và làm "
    "việc xuyên biên giới giữa những người sử dụng các ngôn ngữ khác nhau ngày càng trở nên phổ "
    "biến. Tiếng Anh là ngôn ngữ phổ biến nhất trong giao tiếp quốc tế, tuy nhiên rào cản ngôn "
    "ngữ vẫn là trở ngại lớn đối với nhiều người Việt Nam khi tham gia các buổi hội thảo trực "
    "tuyến, bài giảng quốc tế hoặc giao tiếp trực tiếp. Do đó, nhu cầu về công cụ hỗ trợ giao "
    "tiếp đa ngôn ngữ theo thời gian thực là rất cấp thiết.",

    "Dịch tiếng nói trực tiếp là bài toán khó, đòi hỏi kết hợp ba công nghệ cốt lõi: nhận dạng "
    "tiếng nói, dịch máy và tổng hợp tiếng nói. Thách thức lớn nhất là độ trễ đầu cuối: người "
    "dùng mong muốn thấy phụ đề và nghe bản dịch gần như tức thời, trong khi các hệ thống dịch "
    "máy truyền thống hoạt động theo kiểu hàng loạt, chỉ dịch sau khi câu nói đã kết thúc. Ngoài "
    "ra, chất lượng dịch phụ thuộc vào ngữ cảnh hội thoại, khả năng phân tách người nói và độ "
    "chính xác nhận dạng ở điều kiện tiếng ồn thực tế.",

    "Kết quả khảo sát thực tế các ứng dụng phổ biến (Google Translate, Microsoft Translator, "
    "Saydi) cho thấy: phần lớn chỉ dịch sau khi câu đã hoàn chỉnh nên độ trễ lớn; ít tối ưu riêng "
    "cho cặp ngôn ngữ Anh - Việt; phụ thuộc hoàn toàn vào dịch vụ đám mây nên không thể triển khai "
    "độc lập hoặc dùng trong môi trường yêu cầu bảo mật. Đáng chú ý, một số sản phẩm quảng cáo xử "
    "lý \"trên thiết bị\" nhưng thực tế vẫn gửi dữ liệu âm thanh lên đám mây (qua Deepgram, "
    "Google), gây rủi ro về quyền riêng tư. Vì vậy, việc nghiên cứu và lựa chọn các mô hình mã "
    "nguồn mở để xây dựng hệ thống tự chủ, có khả năng chạy hoàn toàn cục bộ, là vấn đề cần giải "
    "quyết.",

    "Từ những lý do trên, việc nghiên cứu, xây dựng hệ thống dịch tiếng nói trực tiếp Anh - Việt "
    "độ trễ thấp, bảo đảm chất lượng nhận dạng, chất lượng dịch và trải nghiệm người dùng gần với "
    "giao tiếp thời gian thực là cấp thiết cả về lý luận và thực tiễn.",

    "4. Mục tiêu:",
    "- Nghiên cứu, đánh giá và lựa chọn mô hình AI mã nguồn mở cho nhận dạng tiếng nói và dịch "
    "máy Anh - Việt, phù hợp với tài nguyên tính toán thực tế.",
    "- Đề xuất kiến trúc xử lý luồng âm thanh dạng dòng (cascaded kết hợp cơ chế tịnh tiến) nhằm "
    "tối ưu độ trễ, cho phép hiển thị phụ đề ngay trong khi người nói đang nói.",
    "- Xây dựng ứng dụng dịch thuật giọng nói Anh - Việt thời gian thực: thu âm từ micro, hiển "
    "thị phụ đề song ngữ trên hai màn hình tách biệt, tạo phòng dịch cho nhiều người nghe truy "
    "cập bằng mã QR.",
    "- Hỗ trợ hai chế độ triển khai: chạy hoàn toàn cục bộ (không gửi dữ liệu ra ngoài) và chạy "
    "kết hợp GPU từ xa (chỉ gửi văn bản, không gửi âm thanh).",
    "- Đo đạc và đánh giá định lượng: WER của nhận dạng, BLEU/chrF của dịch máy và độ trễ đầu "
    "cuối; so sánh cấu hình CPU và GPU.",

    "5. Nội dung:",
    "Đồ án sử dụng các mô hình mã nguồn mở đã được huấn luyện sẵn (không huấn luyện lại từ đầu), "
    "cụ thể như sau:",
    "- Nhận dạng tiếng nói: OpenAI Whisper thông qua faster-whisper (CTranslate2) - phiên bản "
    "tiny.en (khoảng 39 triệu tham số) cho cấu hình CPU/offline và large-v3-turbo (khoảng 809 "
    "triệu tham số) cho cấu hình có GPU.",
    "- Dịch máy Anh sang Việt: VinAI Translate en2vi-v2 (VinAI Research, công bố tại INTERSPEECH "
    "2022), kiến trúc mBART gồm 12 lớp encoder và 12 lớp decoder, d_model = 1024, 16 head, từ "
    "vựng 66.773 token, khoảng 448 triệu tham số.",
    "- Phát hiện hoạt động tiếng nói: WebRTC VAD kết hợp ngưỡng năng lượng, phân đoạn theo thời "
    "gian thực (đơn vị giây).",
    "- Tổng hợp tiếng nói (tuỳ chọn): Microsoft Edge TTS khi có mạng; Piper/SAPI5 khi chạy "
    "offline.",
    "- Phân tách người nói: đặc trưng MFCC log-mel cài đặt bằng numpy kết hợp phân cụm trực tuyến.",

    "Dữ liệu sử dụng:",
    "- Không huấn luyện lại mô hình; sử dụng trọng số đã huấn luyện sẵn của Whisper và VinAI "
    "Translate.",
    "- Đánh giá dịch máy: bộ câu Anh - Việt chuyên ngành hội họp/công nghệ thông tin kèm bản dịch "
    "tham chiếu do học viên xây dựng, dùng để tính BLEU và chrF.",
    "- Đánh giá nhận dạng: clip âm thanh thật thu từ micro và clip tổng hợp bằng Edge TTS, dùng "
    "để tính WER.",
    "- Đánh giá toàn trình: đo độ trễ p50/p90 trên cùng bộ clip.",
    "- Hướng mở rộng: bổ sung kho ngữ liệu PhoMT (VinAI) và IWSLT2015 (Anh - Việt) để có chỉ số "
    "BLEU so sánh được với các công bố quốc tế.",

    "Nội dung đồ án được chia thành 3 chương, cụ thể như sau:",

    "Chương 1. TỔNG QUAN BÀI TOÁN DỊCH THUẬT GIỌNG NÓI THỜI GIAN THỰC ANH - VIỆT",
    "1.1. Tổng quan nhu cầu giao tiếp đa ngôn ngữ và bài toán dịch tiếng nói trực tiếp",
    "1.2. Nhận dạng tiếng nói, dịch máy và tổng hợp tiếng nói: nguyên lý, công nghệ và các chỉ "
    "số đánh giá (WER, BLEU, chrF, độ trễ)",
    "1.3. Kiến trúc cascaded và simultaneous trong hệ thống dịch tiếng nói thời gian thực",
    "1.4. Khảo sát các hệ thống hiện có (Google Translate, Microsoft Translator, Saydi) và bài "
    "toán đặt ra",
    "1.5. Các mô hình mã nguồn mở được lựa chọn: Whisper, VinAI Translate, WebRTC VAD, "
    "Edge TTS/Piper",

    "Chương 2. PHÂN TÍCH, THIẾT KẾ HỆ THỐNG DỊCH TIẾNG NÓI TRỰC TIẾP ANH - VIỆT",
    "2.1. Mục đích, yêu cầu, phương pháp và phạm vi xây dựng hệ thống",
    "2.2. Kiến trúc tổng quan và luồng xử lý dữ liệu âm thanh của hệ thống",
    "2.3. Thiết kế phân hệ thu âm, phát hiện hoạt động tiếng nói và phân đoạn âm thanh theo "
    "thời gian",
    "2.4. Thiết kế pipeline tịnh tiến: sinh phụ đề từng phần (partial), chốt câu (final) và cơ "
    "chế quản lý độ trễ",
    "2.5. Thiết kế phiên dịch trực tiếp: phòng dịch, mã QR cho người nghe, hiển thị phụ đề song "
    "ngữ trên hai màn hình tách biệt",
    "2.6. Thiết kế hai chế độ triển khai: chạy cục bộ (offline, không gửi dữ liệu) và kết hợp "
    "GPU từ xa (online, chỉ gửi văn bản)",
    "2.7. Thiết kế giao diện ứng dụng",

    "Chương 3. TRIỂN KHAI THỰC NGHIỆM VÀ ĐÁNH GIÁ KẾT QUẢ",
    "3.1. Mô hình hoạt động của hệ thống",
    "3.2. Cài đặt ứng dụng",
    "- Cài đặt môi trường và thư viện mã nguồn mở",
    "- Xây dựng giao diện ứng dụng",
    "- Chuẩn bị mô hình và dữ liệu đánh giá",
    "3.3. Triển khai và kết quả thử nghiệm",
    "- Thử nghiệm dịch hội thoại ngắn hàng ngày",
    "- Thử nghiệm dịch bài giảng, tin tức tiếng Anh",
    "- Thử nghiệm nhiều người nghe đồng thời qua mã QR",
    "3.4. Đánh giá kết quả",
    "- Đo đạc độ trễ toàn trình (p50/p90) trên cấu hình CPU và GPU",
    "- Đánh giá chỉ số WER của nhận dạng và BLEU/chrF của dịch máy",
    "- So sánh mô hình, cấu hình và so sánh với ứng dụng thương mại (Google Translate, "
    "Microsoft Translator, Saydi)",

    "6. Sản phẩm/ kết quả dự kiến:",
    "- Báo cáo đồ án tốt nghiệp, báo cáo tóm tắt.",
    "- Sản phẩm phần mềm: Ứng dụng dịch thuật giọng nói Anh - Việt thời gian thực, cho phép thu "
    "âm trực tiếp từ micro, hiển thị phụ đề song ngữ trên hai màn hình tách biệt (màn hình 1 chỉ "
    "tiếng Anh, màn hình 2 chỉ tiếng Việt), tạo phòng dịch cho nhiều người nghe truy cập bằng mã "
    "QR, xuất phụ đề ra TXT/SRT/VTT, tóm tắt nội dung và phân tách người nói.",
    "- Ứng dụng hỗ trợ hai chế độ: chạy hoàn toàn cục bộ trên máy tính (bảo đảm không gửi dữ "
    "liệu ra ngoài) và chế độ kết hợp GPU miễn phí từ xa để giảm độ trễ (chỉ gửi văn bản).",
    "- Bộ công cụ đo đạc, đánh giá kèm số liệu thực nghiệm về độ trễ, WER và BLEU/chrF.",

    "7. Tiến độ thực hiện: 2,5 tháng (Từ tháng 10 năm 2026 đến tháng 12 năm 2026).",
]


def ptext(p) -> str:
    return "".join(r.text for r in p.runs) or p.text or ""


def main():
    if not os.path.exists(SRC):
        print(f"Khong thay {SRC}")
        return 1
    d = docx.Document(SRC)
    paras = list(d.paragraphs)

    start = end = None
    for i, p in enumerate(paras):
        t = ptext(p).strip()
        if start is None and t.startswith("3. Tính cấp thiết"):
            start = i
        if t.startswith("7. Tiến độ thực hiện"):
            end = i
    if start is None or end is None:
        print(f"Khong tim thay moc: start={start} end={end}")
        return 1
    print(f"Thay the doan van tu chi so {start} den {end} "
          f"({end - start + 1} doan)")

    tmpl = copy.deepcopy(paras[start + 1]._p)     # mau dinh dang doan van thuong
    anchor_el = paras[start - 1]._p               # "2. Thuoc nganh..."
    anchor_parent = paras[start - 1]._parent

    for p in paras[start:end + 1]:
        p._element.getparent().remove(p._element)

    prev = anchor_el
    for text in NOI_DUNG:
        new_el = copy.deepcopy(tmpl)
        prev.addnext(new_el)
        np = Paragraph(new_el, anchor_parent)
        runs = np.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            np.add_run(text)
        prev = new_el

    d.save(DST)
    print(f"Da luu: {DST}")

    # kiem chung lai
    d2 = docx.Document(DST)
    got = [ptext(p).strip() for p in d2.paragraphs if ptext(p).strip()]
    print(f"Tong so doan co chu: {len(got)}")
    for k in ("3. Tính cấp thiết", "Whisper", "VinAI Translate en2vi-v2", "mBART",
              "WebRTC VAD", "PhoMT", "Chương 3", "7. Tiến độ", "Nguyễn Thị Lan"):
        hit = any(k in g for g in got)
        print(f"  {'OK ' if hit else 'THIEU'} {k}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
